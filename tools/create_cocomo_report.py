from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from pathlib import Path

OUT = Path('docs/COCOMO_Cost_Estimation_Report_Salary_Pro.docx')
OUT.parent.mkdir(parents=True, exist_ok=True)

def shade(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), color)
    tc_pr.append(shd)

def set_cell_text(cell, text, bold=False, color=None, align=None):
    cell.text = ''
    p = cell.paragraphs[0]
    if align is not None: p.alignment = align
    r = p.add_run(str(text)); r.bold = bold; r.font.size = Pt(10)
    if color: r.font.color.rgb = RGBColor(*color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = 'Table Grid'
    for i, h in enumerate(headers):
        set_cell_text(t.rows[0].cells[i], h, True, (255,255,255), WD_ALIGN_PARAGRAPH.CENTER)
        shade(t.rows[0].cells[i], '1E3A5F')
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value, False, None, WD_ALIGN_PARAGRAPH.CENTER if i > 0 else None)
            if ri % 2 == 1: shade(cells[i], 'F3F6FA')
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths): row.cells[i].width = Inches(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return t

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(.7); sec.bottom_margin = Inches(.7); sec.left_margin = Inches(.8); sec.right_margin = Inches(.8)

styles = doc.styles
styles['Normal'].font.name = 'Aptos'; styles['Normal'].font.size = Pt(10.5)
styles['Normal'].paragraph_format.space_after = Pt(7)
for style_name, size in [('Title', 22), ('Heading 1', 15), ('Heading 2', 12)]:
    st = styles[style_name]; st.font.name = 'Aptos Display'; st.font.size = Pt(size); st.font.color.rgb = RGBColor(0,0,0)

title = doc.add_paragraph(style='Title'); title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.add_run('COCOMO Cost Estimation Report')
subtitle = doc.add_paragraph(); subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = subtitle.add_run('Salary Pro Salary Management System'); r.bold = True; r.font.size = Pt(13)
doc.add_paragraph()

doc.add_heading('1 Project Overview', level=1)
doc.add_paragraph('Salary Pro is a full stack salary management system developed using React, Tailwind CSS, Node.js, Express, and MongoDB. The system provides user authentication, role based access, employee management, attendance tracking, leave management, salary calculation, payslip generation, reports, settings, and salary insights.')
doc.add_paragraph('This report estimates the software development effort, time, staffing, and cost using the Basic COCOMO model. The estimate is based on the verified source code count from the frontend and backend application files.')

doc.add_heading('2 Project Size Estimation', level=1)
table(doc, ['Source area', 'Files', 'Lines of code'], [
    ['Frontend client src', '23', '381'],
    ['Backend server src', '22', '419'],
    ['Total application source', '45', '800'],
], [3.0, 1.3, 1.8])
doc.add_paragraph('The project size used for COCOMO is calculated as follows:')
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('KLOC = Total LOC / 1000 = 800 / 1000 = 0.8 KLOC'); r.bold = True; r.font.size = Pt(11)
doc.add_paragraph('Third party dependencies, node_modules, lock files, generated files, and diagrams are excluded from the source code count.')

doc.add_heading('3 Chosen COCOMO Project Category', level=1)
doc.add_paragraph('The Organic category is selected because Salary Pro is a relatively small academic web application developed by a small team using familiar technologies and clearly defined business requirements.')
table(doc, ['COCOMO category', 'a', 'b', 'c', 'd'], [['Organic', '2.4', '1.05', '2.5', '0.38']], [2.3, .8, .8, .8, .8])

doc.add_heading('4 COCOMO Estimation Calculations', level=1)
doc.add_heading('4.1 Effort Estimation', level=2)
doc.add_paragraph('Formula: Effort = a x (KLOC)^b')
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('Effort = 2.4 x (0.8)^1.05 = 1.90 person months'); r.bold = True; r.font.size = Pt(11)
doc.add_heading('4.2 Development Time Estimation', level=2)
doc.add_paragraph('Formula: Development Time = c x (Effort)^d')
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('Development Time = 2.5 x (1.90)^0.38 = 3.19 months'); r.bold = True; r.font.size = Pt(11)
doc.add_heading('4.3 Average Team Size', level=2)
doc.add_paragraph('Formula: Average Team Size = Effort / Development Time')
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('Average Team Size = 1.90 / 3.19 = 0.60, approximately 1 person'); r.bold = True; r.font.size = Pt(11)

doc.add_heading('5 Cost Estimation', level=1)
doc.add_paragraph('For this academic estimate, the assumed development cost is Rs 20,000 per person month.')
table(doc, ['Cost item', 'Calculation', 'Result'], [
    ['Estimated effort', '2.4 x (0.8)^1.05', '1.90 person months'],
    ['Cost per person month', 'Assumed rate', 'Rs 20,000'],
    ['Estimated cost', '1.90 x Rs 20,000', 'Rs 38,000'],
], [2.0, 2.5, 1.5])

doc.add_heading('6 Final Estimation Summary', level=1)
table(doc, ['Metric', 'Estimated value'], [
    ['Project size', '800 LOC or 0.8 KLOC'],
    ['Chosen project category', 'Organic'],
    ['Estimated effort', '1.90 person months'],
    ['Estimated development time', '3.19 months'],
    ['Average COCOMO team size', 'Approximately 1 person'],
    ['Actual academic team size', '3 members'],
    ['Assumed cost per person month', 'Rs 20,000'],
    ['Estimated project cost', 'Rs 38,000'],
], [3.2, 2.8])

doc.add_heading('7 Conclusion', level=1)
doc.add_paragraph('Using the Basic COCOMO Organic model, the Salary Pro system with 800 LOC is estimated to require 1.90 person months of effort and 3.19 months of development time. At an assumed cost of Rs 20,000 per person month, the estimated total cost is Rs 38,000. The project was completed by three students who can work in parallel on frontend, backend, database, testing, documentation, and deployment tasks. The COCOMO staffing value is an estimation based on code size and does not replace the actual academic team allocation.')

footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.add_run('Salary Pro COCOMO Cost Estimation Report').font.size = Pt(8)
doc.core_properties.title = 'COCOMO Cost Estimation Report Salary Pro'
doc.core_properties.author = 'Salary Pro Project Team'
doc.save(OUT)
print(OUT.resolve())
