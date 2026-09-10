from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

src=Path('output/Salary_Pro_Software_Requirements_Specification.docx')
dst=Path('output/Salary_Pro_SRS_Traditional_Format.docx')
doc=Document(src)

def set_font(run, size=None, bold=None):
    run.font.name='Times New Roman'
    run._element.rPr.rFonts.set(qn('w:ascii'),'Times New Roman')
    run._element.rPr.rFonts.set(qn('w:hAnsi'),'Times New Roman')
    run.font.color.rgb=RGBColor(0,0,0)
    if size: run.font.size=Pt(size)
    if bold is not None: run.bold=bold
def clear_shading(cell):
    tcPr=cell._tc.get_or_add_tcPr()
    for shd in tcPr.findall(qn('w:shd')): tcPr.remove(shd)
def set_cell_border(cell, **kwargs):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr(); borders = tcPr.first_child_found_in('w:tcBorders')
    if borders is None: borders=OxmlElement('w:tcBorders'); tcPr.append(borders)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        if edge in kwargs:
            tag='w:{}'.format(edge); element=borders.find(qn(tag))
            if element is None: element=OxmlElement(tag); borders.append(element)
            for key,value in kwargs[edge].items(): element.set(qn('w:{}'.format(key)),str(value))

sec=doc.sections[0]
sec.top_margin=Inches(.85); sec.bottom_margin=Inches(.7); sec.left_margin=Inches(.85); sec.right_margin=Inches(.85)
for style_name, size, bold in [('Normal',12,False),('Title',16,True),('Heading 1',14,True),('Heading 2',12,True),('Heading 3',12,True)]:
    s=doc.styles[style_name]; s.font.name='Times New Roman'; s._element.rPr.rFonts.set(qn('w:ascii'),'Times New Roman'); s._element.rPr.rFonts.set(qn('w:hAnsi'),'Times New Roman'); s.font.size=Pt(size); s.font.color.rgb=RGBColor(0,0,0); s.font.bold=bold

for para in doc.paragraphs:
    name=para.style.name
    para.paragraph_format.space_after=Pt(5)
    if name.startswith('Heading'):
        para.alignment=WD_ALIGN_PARAGRAPH.LEFT
        para.paragraph_format.space_before=Pt(8)
    elif name=='Title': para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    elif name.startswith('List Bullet'): para.alignment=WD_ALIGN_PARAGRAPH.LEFT
    else: para.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    for run in para.runs: set_font(run, 16 if name=='Title' else (14 if name=='Heading 1' else (12 if name.startswith('Heading') else 12)))
for table in doc.tables:
    for row_i,row in enumerate(table.rows):
        for cell in row.cells:
            clear_shading(cell); set_cell_border(cell,top={'val':'single','sz':'4','color':'000000'},bottom={'val':'single','sz':'4','color':'000000'},left={'val':'single','sz':'4','color':'000000'},right={'val':'single','sz':'4','color':'000000'})
            for para in cell.paragraphs:
                para.alignment=WD_ALIGN_PARAGRAPH.LEFT
                para.paragraph_format.space_after=Pt(0)
                for run in para.runs: set_font(run,10,bold=(row_i==0))
for para in sec.header.paragraphs+sec.footer.paragraphs:
    para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for run in para.runs: set_font(run,9)
doc.save(dst)
print(dst.resolve())
