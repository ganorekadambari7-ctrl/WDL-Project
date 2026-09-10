from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from pathlib import Path
from datetime import date

OUT = Path('output')
OUT.mkdir(exist_ok=True)
DEST = OUT / 'Salary_Pro_Software_Requirements_Specification.docx'

doc = Document()
sec = doc.sections[0]
sec.top_margin = sec.bottom_margin = Inches(0.8)
sec.left_margin = sec.right_margin = Inches(0.85)

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); tcPr.append(shd)
def cell_text(cell, value, bold=False, color=None):
    cell.text = ''
    p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
    r = p.add_run(str(value)); r.bold = bold; r.font.size = Pt(9)
    if color: r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers)); t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i,h in enumerate(headers):
        cell_text(t.rows[0].cells[i],h,True,'FFFFFF'); shade(t.rows[0].cells[i],'4C1D95')
    for row in rows:
        cells=t.add_row().cells
        for i,v in enumerate(row): cell_text(cells[i],v)
    if widths:
        for row in t.rows:
            for i,w in enumerate(widths): row.cells[i].width=Inches(w)
    doc.add_paragraph()
    return t
def p(text='', style=None, bold_lead=None):
    para=doc.add_paragraph(style=style)
    para.paragraph_format.space_after=Pt(6)
    if bold_lead and text.startswith(bold_lead):
        r=para.add_run(bold_lead); r.bold=True; para.add_run(text[len(bold_lead):])
    else: para.add_run(text)
    return para
def bullet(text, level=0):
    para=doc.add_paragraph(style='List Bullet' if level==0 else 'List Bullet 2'); para.paragraph_format.space_after=Pt(3); para.add_run(text); return para
def req(rid, text, priority='High', status='Implemented'):
    table(['ID','Requirement','Priority','Status'], [[rid,text,priority,status]], [0.8,4.55,0.7,0.95])
def heading(text, level=1): doc.add_heading(text,level=level)
def pagebreak(): doc.add_page_break()

# Styles
styles=doc.styles
styles['Normal'].font.name='Calibri'; styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'),'Calibri'); styles['Normal'].font.size=Pt(10.5)
styles['Normal'].paragraph_format.line_spacing=1.15
for name,size,color in [('Title',28,'2E1065'),('Heading 1',17,'4C1D95'),('Heading 2',13,'5B21B6'),('Heading 3',11,'4C1D95')]:
    st=styles[name]; st.font.name='Calibri'; st.font.size=Pt(size); st.font.color.rgb=RGBColor.from_string(color); st.font.bold=True
    st.paragraph_format.space_before=Pt(14); st.paragraph_format.space_after=Pt(6)

# Header/footer
header=sec.header.paragraphs[0]; header.text='SALARY PRO | SOFTWARE REQUIREMENTS SPECIFICATION'; header.runs[0].font.size=Pt(8); header.runs[0].font.color.rgb=RGBColor(100,116,139)
footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
footer.add_run('Salary Pro - Version 1.0 | Confidential Academic Project') .font.size=Pt(8)

# Cover
doc.add_paragraph().paragraph_format.space_after=Pt(40)
q=doc.add_paragraph(); q.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=q.add_run('SOFTWARE REQUIREMENTS\nSPECIFICATION'); r.bold=True; r.font.size=Pt(29); r.font.color.rgb=RGBColor(76,29,149)
q=doc.add_paragraph(); q.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=q.add_run('Salary Pro'); r.bold=True; r.font.size=Pt(25); r.font.color.rgb=RGBColor(15,23,42)
q=doc.add_paragraph(); q.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=q.add_run('Salary Management System'); r.font.size=Pt(14); r.font.color.rgb=RGBColor(71,85,105)
doc.add_paragraph().paragraph_format.space_after=Pt(30)
table(['Document information','Value'],[
 ['Version','1.0 - baseline requirements'],['Status','Draft for academic review'],['Prepared for','Salary Pro project team and academic evaluators'],['Prepared on',date.today().strftime('%d %B %Y')],['Technology baseline','React, Tailwind CSS, Express.js, MongoDB / Mongoose'],['Document basis','Current project source code and the supplied SRS structure']
],[2.1,4.4])
p('Purpose statement', 'Heading 2'); p('This document defines the current and planned requirements for Salary Pro. Implemented items are marked “Implemented”; intentionally forward-looking requirements are marked “Planned” so that the specification remains an honest development roadmap.')
pagebreak()

heading('Revision History')
table(['Version','Date','Author','Summary'],[['1.0',date.today().strftime('%d %b %Y'),'Salary Pro Project Team','Initial SRS based on the implemented repository and planned extensions.']],[0.9,1.2,1.8,2.6])
heading('Table of Contents')
for x in ['1. Introduction','2. Overall Description','3. External Interface Requirements','4. System Features','5. Other Nonfunctional Requirements','6. Other Requirements','Appendix A: Glossary','Appendix B: Analysis Model','Appendix C: To Be Determined List']:
    p(x)
pagebreak()

heading('1. Introduction')
heading('1.1 Purpose',2)
p('Salary Pro is a web-based salary and workforce-management application for organizations that need a single workspace to maintain employee records, record attendance, manage leave requests, calculate payroll, generate payslips, and review salary insights. This SRS establishes functional, interface, quality, security, and future-scope requirements for the system.')
heading('1.2 Document Conventions',2)
bullet('Numbered sections follow the IEEE 830-style SRS structure demonstrated in the reference document.')
bullet('Requirements use unique identifiers: UI for user interface, SI for software interface, FR for functional requirement, NFR for nonfunctional requirement, BR for business rule, and OR for other requirement.')
bullet('Priority values are High, Medium, or Low. “Implemented” denotes a feature present in the reviewed repository; “Planned” denotes a future requirement.')
bullet('Terms such as JWT, API, and payslip are defined in Appendix A.')
heading('1.3 Intended Audience and Reading Suggestions',2)
table(['Audience','Primary use of this document'],[
 ['Project team','Use Sections 3-5 to implement, test, and plan releases.'],['Academic evaluators','Use Sections 1, 2, and 4 to assess scope and requirements traceability.'],['System administrators','Use Sections 3, 5, and 6 for deployment, security, and operational expectations.'],['Future maintainers','Use the interfaces, data model, and TBD list to safely extend the application.']
],[1.55,4.95])
heading('1.4 Product Scope',2)
p('Salary Pro consolidates routine payroll operations that are often maintained in separate spreadsheets. The system supports authorized administrators and employees through a responsive browser application. Its current core covers identity and access management, employee directory management, attendance, leave, salary calculation, downloadable payslips, reports, insights, UML documentation, and account settings.')
p('Objectives:',bold_lead='Objectives:')
for x in ['Reduce manual salary calculation errors through a consistent net-salary formula.','Provide an auditable record of employee, attendance, leave, and salary data.','Give administrators a simple approval workflow and employees a clear self-service experience.','Provide actionable salary and attendance indicators without requiring a separate analytics product.','Establish a modular foundation for statutory payroll, reporting, notifications, and integrations.']: bullet(x)
p('Current exclusions:',bold_lead='Current exclusions:')
for x in ['Statutory tax calculation rules, PF/ESI computation, and government filing submission.','Bank transfer initiation or integration with payment gateways.','Biometric attendance devices and geofencing.','Native Android or iOS applications.','Automatic AI/LLM analysis; current insights are deterministic calculations from stored records.']: bullet(x)
heading('1.5 References',2)
for x in ['IEEE Std 830-1998, Recommended Practice for Software Requirements Specifications.','Salary Pro source repository: React client and Express/MongoDB server reviewed for this document.','MongoDB and Mongoose documentation for persistence implementation.','React, Vite, Tailwind CSS, Express, JSON Web Token, bcryptjs, and jsPDF documentation.','Supplied “Sample SRS Document.pdf” for the requested document organization.']: bullet(x)
pagebreak()

heading('2. Overall Description')
heading('2.1 Product Perspective',2)
p('Salary Pro is a standalone full-stack web application. The React single-page frontend sends JSON requests to an Express REST API. The API authenticates users using signed JSON Web Tokens, applies role checks where required, and stores persistent operational data in MongoDB through Mongoose models. The application is designed as a new academic product and can later integrate with attendance hardware, statutory-payroll services, notification providers, and banking systems.')
heading('2.2 Product Functions',2)
table(['Function','Current capability'],[
 ['Authentication and roles','Register and sign in as admin or employee; store an authenticated browser session.'],['Employee directory','Create, view, search, edit, and delete employee records.'],['Attendance','Record present/absent status by employee and date; show monthly summary.'],['Leave management','Create leave requests; administrators approve or reject pending requests.'],['Payroll','Calculate and save basic salary, bonus, tax deduction, late-mark deduction, total deductions, and net salary.'],['Payslips','Download a formatted PDF payslip from saved salary data.'],['Reports and insights','Show dashboard/report visualizations and calculate salary anomalies, trends, and suggestions.'],['Settings','Update profile, password, profile image, and, for admins, user roles.']
],[1.5,5.0])
heading('2.3 User Classes and Characteristics',2)
table(['Class','Characteristics and permissions'],[
 ['Administrator','Manages employee directory, attendance, payroll, account roles, and leave approvals. Familiar with HR/payroll operations.'],['Employee','Registers/logs in, accesses protected workspace pages, submits leave requests, and views operational information.'],['Project maintainer','Configures deployment and database, resolves defects, and develops planned modules.']
],[1.35,5.15])
heading('2.4 Operating Environment',2)
for x in ['Client: modern Chrome, Edge, Firefox, or mobile browser; responsive React 18 application built with Vite and Tailwind CSS.','Server: Node.js runtime with Express 4 API, JSON parsing, CORS policy, JWT, bcryptjs, and Mongoose.','Data store: MongoDB accessible through a configured MONGODB_URI connection string.','Deployment: separate static frontend and Node/Express API hosting are expected; HTTPS is required in production.','Optional client document generation: jsPDF creates a downloadable A4 payslip in the user browser.']: bullet(x)
heading('2.5 Design and Implementation Constraints',2)
for x in ['The current client and server are JavaScript/ES module applications using the project’s specified stack.','Passwords must never be returned by the API and are stored using bcrypt hashing.','Employee IDs are unique identifiers and must be normalized consistently.','Attendance permits one record per employee per date through a unique compound index.','The present system uses stored records and rule-based insights; it must not claim predictive AI analysis without a configured model or provider.','Current dashboard/report example figures are presentation data in some views and must be connected to live data before production use.']: bullet(x)
heading('2.6 User Documentation',2)
p('The delivered documentation set shall include this SRS, a setup guide (README), API route documentation, a user guide for employee, attendance, leave, payroll, and settings tasks, and a deployment/environment-variable guide. Contextual validation and notices in the UI shall help users correct errors.')
heading('2.7 Assumptions and Dependencies',2)
for x in ['A MongoDB instance and valid MONGODB_URI are available to the server.','A strong JWT_SECRET and client URL are configured as environment variables outside source control.','Users access the app through a supported browser and retain local browser storage for the current login session.','The organization supplies accurate employee, attendance, leave, and deduction inputs.','Future statutory and banking functions depend on approved regional rules and external-provider contracts.']: bullet(x)
pagebreak()

heading('3. External Interface Requirements')
heading('3.1 User Interfaces',2)
table(['ID','Interface requirement','Status'],[
 ['UI-1','Login and registration pages shall collect role, identity, and password and return understandable validation errors.','Implemented'],['UI-2','The protected application shell shall provide sidebar navigation, top bar, account menu, and sign-out access.','Implemented'],['UI-3','Employee management shall provide a directory, search, and add/edit/delete form with validation.','Implemented'],['UI-4','Attendance shall support employee/date/status entry and monthly per-employee summaries.','Implemented'],['UI-5','Leave management shall provide a request form, status display, and admin approval/rejection controls.','Implemented'],['UI-6','Payroll shall show salary components and provide a downloadable payslip action.','Implemented'],['UI-7','Dashboard, reports, and insights shall distinguish live metrics from sample/demo values.','Planned refinement'],['UI-8','The interface shall be responsive for desktop and small mobile screens and provide labels for interactive controls.','Implemented / ongoing']
],[0.55,4.95,1.0])
heading('3.2 Hardware Interfaces',2)
p('HI-1: No direct hardware interface is required in the current release. The system operates through standard desktop and mobile browsers. Future biometric device or card-reader interfaces are outside the current baseline.')
heading('3.3 Software Interfaces',2)
table(['ID','Interface','Requirement'],[
 ['SI-1','React frontend','React Router shall manage protected navigation; the AuthContext shall keep the authenticated user and token state in browser local storage.'],['SI-2','Express REST API','The backend shall expose JSON endpoints for auth, employees, attendance, leave, salary, settings, insights, and health checks.'],['SI-3','MongoDB / Mongoose','Persistent models shall represent User, Employee, Attendance, Leave, and Salary records.'],['SI-4','JWT and bcryptjs','JWT shall identify authenticated requests; bcryptjs shall hash and verify passwords.'],['SI-5','jsPDF','The browser client shall generate A4 salary slips using saved employee and salary data.'],['SI-6','Future integrations','Statutory payroll, notifications, bank transfers, and biometric attendance require separately approved APIs and credentials.']
],[0.55,1.5,4.45])
heading('3.4 Communications Interfaces',2)
for x in ['CI-1: Client-server communication shall use JSON over HTTP during local development and HTTPS in production.','CI-2: Authenticated API calls shall send the JWT in an Authorization: Bearer header.','CI-3: CORS shall allow only the configured client URL(s); unapproved origins shall be rejected.','CI-4: MongoDB connectivity shall use the configured secure database connection string.','CI-5: Future third-party integrations shall use HTTPS, least-privilege credentials, timeout handling, retries where appropriate, and audit logs.']: bullet(x)
pagebreak()

heading('4. System Features')
heading('4.1 Authentication and Access Control',2)
p('Description and priority: High. Salary Pro protects its workspace with role-aware registration and login. Admin and employee are the two current roles.')
p('Stimulus/response: A visitor submits valid registration details; the system creates a bcrypt-hashed user and returns a seven-day JWT. A user submits valid credentials with the selected role; the system returns a session. An invalid request returns a safe error message. A protected request without a valid token receives a 401 response.')
req('FR-AUTH-1','The system shall require name, username, email, password, and role for registration and accept only admin or employee roles.','High','Implemented')
req('FR-AUTH-2','The system shall validate username, email-domain, and minimum password rules before creating an account.','High','Implemented')
req('FR-AUTH-3','The system shall hash passwords with bcrypt before persistence and exclude password hashes from normal API responses.','High','Implemented')
req('FR-AUTH-4','The system shall authenticate protected API requests with a verifiable JWT and return 401 for invalid or expired sessions.','High','Implemented')
req('FR-AUTH-5','The system shall restrict leave review, user listing, and role changes to administrators.','High','Implemented')
req('FR-AUTH-6','The future system shall implement password-reset and account-lockout workflows.','Medium','Planned')
heading('4.2 Employee Directory',2)
p('Description and priority: High. Administrators maintain a canonical employee directory that payroll, attendance, and leave functions reference through employee ID.')
req('FR-EMP-1','The system shall create employee records containing name, unique employee ID, date of birth, education, experience, and basic salary.','High','Implemented')
req('FR-EMP-2','The system shall list employees in descending creation order and allow client-side searching by name or employee ID.','High','Implemented')
req('FR-EMP-3','The system shall permit update and deletion of an existing employee record and return suitable not-found/conflict responses.','High','Implemented')
req('FR-EMP-4','The system shall prevent duplicate employee IDs and negative experience or salary values.','High','Implemented')
req('FR-EMP-5','The future system shall preserve historical payroll records and warn before deleting employees with related records.','High','Planned')
heading('4.3 Attendance Management',2)
p('Description and priority: High. Attendance records support payroll review and attendance-rate insights.')
req('FR-ATT-1','The system shall record present or absent attendance for an employee on a specified date.','High','Implemented')
req('FR-ATT-2','The system shall upsert a record for the employee/date pair so only one record exists per employee per calendar date.','High','Implemented')
req('FR-ATT-3','The system shall retrieve an employee’s monthly attendance records and calculate total, present, absent, and percentage values.','High','Implemented')
req('FR-ATT-4','The future system shall support late, half-day, work-from-home, holiday, and approved-leave statuses.','Medium','Planned')
pagebreak()

heading('4. System Features (continued)')
heading('4.4 Leave Management',2)
p('Description and priority: High. Employees can submit requests and administrators can review pending requests.')
req('FR-LEV-1','The system shall collect employee ID, start/end date, leave type, reason, and optional late-mark flag for a leave request.','High','Implemented')
req('FR-LEV-2','The system shall reject a leave request whose end date precedes its start date.','High','Implemented')
req('FR-LEV-3','The system shall assign pending status on creation and list leave records newest first.','High','Implemented')
req('FR-LEV-4','Only an administrator shall approve or reject a pending leave request; the reviewer ID shall be retained.','High','Implemented')
req('FR-LEV-5','The future system shall calculate leave balances, detect overlapping requests, and notify requesters of decisions.','Medium','Planned')
heading('4.5 Payroll Calculation and Payslips',2)
p('Description and priority: Critical. Payroll records calculate transparent net salary values and provide a user-downloadable payslip.')
p('Net salary business formula: Net Salary = max(0, Basic Salary + Bonus - (Tax Deduction + Late-Mark Deduction)).')
req('FR-PAY-1','The system shall accept a selected employee ID, basic salary, bonus, tax deduction, and late-mark deduction, all as non-negative values.','High','Implemented')
req('FR-PAY-2','The server shall calculate total deductions and net salary and persist the resulting salary record.','High','Implemented')
req('FR-PAY-3','The system shall retrieve salary records for a given employee ID in newest-first order.','High','Implemented')
req('FR-PAY-4','The client shall generate an A4 PDF payslip that includes employee identity, payroll period, earnings, deductions, net salary, and generation date.','High','Implemented')
req('FR-PAY-5','The future system shall introduce explicit payroll periods, draft/review/approved/paid states, bulk pay runs, and immutable approval history.','High','Planned')
req('FR-PAY-6','The future system shall calculate statutory deductions according to approved jurisdictional configurations, not hard-coded assumptions.','High','Planned')
heading('4.6 Reporting and Insights',2)
p('Description and priority: Medium. Current insights derive from saved salaries and attendance; they should be interpreted as rule-based decision support, not as autonomous AI predictions.')
req('FR-INS-1','The system shall calculate latest salary per employee, average net salary, attendance rate, and salary direction from saved records.','Medium','Implemented')
req('FR-INS-2','The system shall flag latest salaries more than 35% above or below the calculated average as anomalies.','Medium','Implemented')
req('FR-INS-3','The system shall offer rule-based suggestions concerning bonus coverage, attendance, and salary outliers.','Medium','Implemented')
req('FR-INS-4','The future system shall make dashboard and reports pages data-backed, filterable by period/department, and exportable.','High','Planned')
req('FR-INS-5','The future system may add natural-language analysis only after secure provider integration, data-governance review, and human oversight are approved.','Low','Planned')
heading('4.7 Profile and Workspace Administration',2)
req('FR-SET-1','The system shall let an authenticated user update profile information including name, contact details, work details, address, and profile image.','Medium','Implemented')
req('FR-SET-2','The system shall validate a ten-digit phone number and allow only JPG/JPEG/PNG images below 2 MB in the client interface.','Medium','Implemented')
req('FR-SET-3','The system shall require current password verification and a minimum eight-character new password when changing password.','High','Implemented')
req('FR-SET-4','Administrators shall list workspace users and change their role between employee and admin.','High','Implemented')
req('FR-SET-5','The future system shall store profile media in managed object storage rather than persisting large data URLs in user records.','Medium','Planned')
pagebreak()

heading('5. Other Nonfunctional Requirements')
heading('5.1 Performance Requirements',2)
table(['ID','Requirement'],[
 ['NFR-PERF-1','The health endpoint shall respond within 2 seconds under normal development/demo load.'],['NFR-PERF-2','Common list and form interactions should complete within 3 seconds after a healthy API/database response.'],['NFR-PERF-3','A single payroll calculation shall complete and return its saved record within 3 seconds under normal load.'],['NFR-PERF-4','The future production deployment shall support at least 100 concurrent active users after load testing.'],['NFR-PERF-5','PDF payslip generation should begin within 5 seconds on a typical current browser.']
],[1.2,5.3])
heading('5.2 Safety Requirements',2)
for x in ['SAF-1: The system shall display monetary values with clear Indian Rupee formatting in relevant UI and payslip displays.','SAF-2: Net salary must never be persisted as negative; the current formula clamps it to zero.','SAF-3: Future salary records must be reviewed/approved before marking them paid.','SAF-4: The product shall state that users must verify statutory deductions and payroll outputs against the organization’s approved policies before payment.','SAF-5: Destructive employee deletion shall require confirmation and, in future, be replaced with safe archival where related records exist.']: bullet(x)
heading('5.3 Security Requirements',2)
for x in ['SEC-1: Passwords shall be hashed with bcrypt and excluded from ordinary user responses.','SEC-2: Protected routes shall verify JWT signatures and reject missing, invalid, or expired tokens.','SEC-3: Administrator-only actions shall enforce server-side authorization, not merely hide UI controls.','SEC-4: Production traffic shall use HTTPS; CORS shall restrict allowed browser origins.','SEC-5: Secrets such as MONGODB_URI and JWT_SECRET shall remain in environment variables and out of source control.','SEC-6: Planned: apply rate limiting, secure HTTP headers, token revocation/rotation, audit logging, least-privilege database users, data retention controls, and an OWASP security review before production launch.']: bullet(x)
heading('5.4 Software Quality Attributes',2)
table(['ID','Quality requirement'],[
 ['NFR-USE-1','Usability: workflows shall provide labeled fields, clear feedback, and responsive layout without special training.'],['NFR-REL-1','Reliability: API errors shall return consistent JSON messages; persistent model validation shall prevent invalid core data.'],['NFR-MNT-1','Maintainability: frontend pages/components and backend models/controllers/routes shall remain modular and documented.'],['NFR-PORT-1','Portability: client shall work on current evergreen desktop and mobile browsers.'],['NFR-ACC-1','Accessibility: forms shall have labels, focus-visible controls, readable contrast, and keyboard-accessible basic actions; WCAG 2.1 AA review is planned.'],['NFR-TST-1','Testability: future releases shall add automated unit, API, and end-to-end tests for critical routes and payroll calculations.']
],[1.35,5.15])
heading('5.5 Business Rules',2)
for x in ['BR-1: Only the roles admin and employee are valid in the current release.','BR-2: An employee ID is unique and is normalized in attendance/salary operations.','BR-3: Attendance status is present or absent in the current data model.','BR-4: Leave status transitions are pending to approved or rejected; only admins may perform review.','BR-5: Total deductions equal tax deduction plus late-mark deduction.','BR-6: Net salary cannot be below zero.','BR-7: The salary-insight anomaly threshold is more than 35% above or below the current average net salary.','BR-8: Future legal payroll rules must be configurable by jurisdiction, employer policy, and effective date.']: bullet(x)
pagebreak()

heading('6. Other Requirements')
for x in ['OR-1: The project shall include source code, this SRS, setup instructions, and a live or recorded demonstration for academic evaluation.','OR-2: A server environment example shall identify required variables without exposing secret values.','OR-3: The project shall maintain API endpoint documentation, data-model descriptions, and release notes as the system evolves.','OR-4: The initial interface language is English; internationalization and localized currency/tax policy support are planned future work.','OR-5: Before handling real organization payroll in production, the project shall complete a security, privacy, backup, recovery, and legal-compliance review.','OR-6: Production deployment shall establish backups, monitoring, database migration strategy, incident response, and user-support ownership.']: bullet(x)
heading('Future Release Roadmap',2)
table(['Release','Planned scope'],[
 ['R1 - Data integrity','Employee archive, historical relationship protection, live dashboard/report data, filters, export, and automated tests.'],['R2 - Payroll operations','Pay periods, bulk runs, approval states, leave-balance calculation, notifications, attendance expansion, audit trail.'],['R3 - Compliance and integration','Configurable statutory payroll, bank/payment export, biometric/provider integration, managed media storage, operational monitoring.'],['R4 - Intelligence','Governed natural-language assistant, forecasting only with approved data policies, evaluation, explainability, and human review.']
],[1.45,5.05])
pagebreak()

heading('Appendix A: Glossary')
table(['Term','Definition'],[
 ['API','Application Programming Interface; the HTTP interface used by the client and server.'],['Admin','A user role permitted to manage users and approve/reject leave requests.'],['Attendance','A dated employee record with current values present or absent.'],['bcrypt','Password-hashing library used to store passwords securely.'],['CORS','Browser policy limiting which web origins may call the API.'],['Employee ID','Unique organization identifier used to associate employee, attendance, leave, and salary records.'],['JWT','JSON Web Token; signed token used to identify an authenticated request.'],['Leave request','Request recording leave dates, duration, reason, late-mark flag, and review status.'],['Mongoose','Object data modeling library used with MongoDB.'],['Net salary','Basic salary plus bonus minus total deductions, not less than zero.'],['Payslip','Employee-facing PDF showing payroll components for a salary record.'],['Salary insight','Rule-based summary or recommendation calculated from stored records.']
],[1.4,5.1])
heading('Appendix B: Analysis Model',1)
heading('B.1 Logical Data Flow',2)
p('User browser -> React application -> authenticated REST API -> Express controllers -> Mongoose models -> MongoDB. The API returns JSON; the React app renders protected pages and may produce a browser-side PDF payslip.')
heading('B.2 Core Entity Relationships',2)
table(['Entity','Key fields','Relationship / purpose'],[
 ['User','name, username, email, password, role','Identity and authorization; a reviewer reference can be stored on Leave.'],['Employee','name, employeeId, education, experience, basicSalary','Canonical business employee record.'],['Attendance','employeeId, date, status','Many attendance records per employee; one per employee/date.'],['Leave','employeeId, dates, leaveType, reason, status, reviewedBy','Many leave requests per employee; reviewedBy references User.'],['Salary','employeeId, basicSalary, bonus, deductions, netSalary','Many salary records per employee; supports payroll history.']
],[1.1,2.65,2.75])
heading('B.3 Requirement Traceability Summary',2)
table(['Area','Primary requirement IDs'],[
 ['Identity & access','FR-AUTH-1 through FR-AUTH-6; SEC-1 through SEC-6'],['Employee & time','FR-EMP-1 through FR-EMP-5; FR-ATT-1 through FR-ATT-4; FR-LEV-1 through FR-LEV-5'],['Payroll','FR-PAY-1 through FR-PAY-6; BR-5 and BR-6; SAF-1 through SAF-4'],['Analytics & reporting','FR-INS-1 through FR-INS-5; NFR-PERF-1 through NFR-PERF-5'],['Administration','FR-SET-1 through FR-SET-5; OR-1 through OR-6']
],[2.0,4.5])
pagebreak()

heading('Appendix C: To Be Determined List')
table(['ID','Open item','Decision needed'],[
 ['TBD-1','Target deployment platform and domain','Choose hosting for React client, Express API, MongoDB region, DNS, TLS, and monitoring.'],['TBD-2','Payroll jurisdiction and statutory rules','Confirm applicable tax, PF/ESI, professional-tax, overtime, and leave policy requirements.'],['TBD-3','Role model expansion','Determine whether HR, payroll approver, manager, finance, and auditor roles are needed.'],['TBD-4','Approval workflow','Confirm payroll-period ownership, approval steps, cut-off dates, correction/reversal policy, and audit retention.'],['TBD-5','Dashboard/report data contract','Replace static/demo visual values with audited live calculations and specify filter/export behavior.'],['TBD-6','Privacy and retention','Define profile-document/media storage, access, retention, deletion, backups, and incident-response policy.'],['TBD-7','Notification channel','Choose email, in-app, SMS, or WhatsApp provider and consent requirements for leave/payroll notifications.'],['TBD-8','External integrations','Approve any biometric attendance, accounting, banking, or AI provider and data-sharing agreement before implementation.'],['TBD-9','Test and acceptance plan','Set test coverage goals, test data policy, performance target, accessibility test method, and release criteria.']
],[0.65,2.1,3.75])
p('End of document')

doc.core_properties.title='Salary Pro Software Requirements Specification'
doc.core_properties.subject='Current and planned requirements for Salary Pro'
doc.core_properties.author='Salary Pro Project Team'
doc.save(DEST)
print(DEST.resolve())
